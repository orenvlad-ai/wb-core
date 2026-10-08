"""Actual Chromium/HTTP FF form: one mutation, same-request recovery and restart."""
import json
from io import BytesIO
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from playwright.sync_api import sync_playwright, expect
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from apps.operator_fulfillment_services_integration_smoke import server_for,stop,UPLOADS,request
from apps.operator_fulfillment_services_native_smoke import seed
from apps.sheet_vitrina_v1_fulfillment_services_smoke import _build_workbook,_valid_row,_wb_supply_row
from packages.adapters.registry_upload_http_entrypoint import DEFAULT_SHEET_OPERATOR_UI_PATH
from packages.application import operator_fulfillment_services as receipt
from packages.application.fulfillment_services import UPLOADS_TABLE


def open_form(page,base):
    page.goto(base+DEFAULT_SHEET_OPERATOR_UI_PATH+'?embedded_tab=factory-order',wait_until='domcontentloaded')
    page.get_by_role('button',name='ФФ',exact=True).click()
    page.get_by_role('button',name='Услуги ФФ',exact=True).click()
    expect(page.locator('#fulfillmentServicesTitle')).to_be_visible()


def choose_workbook(page, file):
    # Use the operator's actual button/chooser. Direct assignment to the hidden
    # input bypasses the busy guard while a previous acceptance/list refresh is
    # still finishing, and can silently discard a synthetic change event.
    with page.expect_file_chooser() as chosen:
        page.locator('#fulfillmentUploadButton').click()
    chosen.value.set_files(file)


def test_chartsheet_diagnostic_retry():
    """A saved rejection ends one action; correcting the file starts a new one."""
    workbook=Workbook();chart=workbook.create_chartsheet();chart.add_chart(BarChart())
    workbook.remove(workbook.active)
    buffer=BytesIO();workbook.save(buffer);chart_data=buffer.getvalue()
    assert not load_workbook(BytesIO(chart_data)).worksheets
    with TemporaryDirectory(prefix='ffsvc-browser-chartsheet-') as raw:
        rt,_=seed(raw);entry,server,thread,base=server_for(rt)
        try:
            with sync_playwright() as pw:
                browser=pw.chromium.launch();page=browser.new_page()
                mutations=[]
                def lose_response(route):
                    if route.request.method=='POST':
                        mutations.append('upload');response=route.fetch()
                        assert response.status==200
                        route.abort('failed')
                    else:route.continue_()
                page.route('**/fulfillment-services/uploads',lose_response)
                open_form(page,base)
                choose_workbook(page, {'name':'chart-only.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':chart_data})
                expect(page.locator('#fulfillmentServicesMessage')).to_contain_text('Документ не принят.',timeout=10000)
                expect(page.locator('#fulfillmentAcceptance')).to_be_hidden()
                expect(page.locator('#fulfillmentUploadButton')).to_be_enabled()
                assert page.locator('#fulfillmentAcceptance .ff-operation-check').count()==0
                with receipt.readonly(rt.db_path) as conn:
                    aliases=[dict(row) for row in conn.execute('SELECT * FROM '+receipt.REQUESTS)]
                    assert len(aliases)==1 and aliases[0]['operation_id'] is None
                    rejected_id=aliases[0]['request_id']
                    failed=conn.execute('SELECT validation_status, validation_error_summary FROM '+UPLOADS_TABLE).fetchone()
                    assert failed['validation_status']=='failed' and 'does not contain a worksheet' in failed['validation_error_summary']
                status,recovered=request(base,UPLOADS+'?request_id='+rejected_id)
                assert status==200 and recovered['status']=='rejected' and recovered['acceptance'] is None
                open_form(page,base)
                expect(page.locator('#fulfillmentUploadButton')).to_be_enabled()
                expect(page.locator('#fulfillmentAcceptance')).to_be_hidden()
                assert mutations==['upload']
                choose_workbook(page, {'name':'corrected.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':_build_workbook([_valid_row('1001')])})
                expect(page.locator('#fulfillmentAcceptance .ff-operation-check')).to_be_visible(timeout=10000)
                expect(page.locator('#fulfillmentAcceptance')).to_contain_text('Документ сохранён.')
                expect(page.locator('#fulfillmentUploadButton')).to_be_enabled()
                with receipt.readonly(rt.db_path) as conn:
                    aliases=[dict(row) for row in conn.execute('SELECT * FROM '+receipt.REQUESTS)]
                    assert len(aliases)==2 and len({row['request_id'] for row in aliases})==2
                    assert next(row for row in aliases if row['operation_id'])['request_id']!=rejected_id
                    assert conn.execute('SELECT count(*) FROM '+UPLOADS_TABLE).fetchone()[0]==2
                    assert conn.execute('SELECT count(*) FROM '+UPLOADS_TABLE+" WHERE validation_status='failed'").fetchone()[0]==1
                    assert conn.execute('SELECT count(*) FROM '+receipt.TABLE).fetchone()[0]==1
                open_form(page,base)
                assert mutations==['upload','upload']
                browser.close()
            print('Chromium chartsheet-only saved diagnostic, lost-response rejected GET, reload no duplicate, corrected file new identity: OK')
        finally:stop(server,thread)


def main():
    with TemporaryDirectory(prefix='ffsvc-browser-') as raw:
        rt,_=seed(raw);entry,server,thread,base=server_for(rt)
        try:
            with sync_playwright() as pw:
                browser=pw.chromium.launch();page=browser.new_page(viewport={'width':1440,'height':1000})
                mutations=[];foreign=True
                def upload(route):
                    if route.request.method=='POST':
                        mutations.append('upload');route.fetch();route.abort('failed')
                    else:route.continue_()
                def recovery(route):
                    response=route.fetch();payload=response.json()
                    if foreign and payload.get('acceptance'):
                        payload['request_id']='another-request'
                    route.fulfill(status=response.status,content_type='application/json',body=json.dumps(payload))
                page.route('**/fulfillment-services/uploads',upload)
                page.route('**/fulfillment-services/uploads?request_id=*',recovery)
                open_form(page,base)
                choose_workbook(page, {'name':'synthetic.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':_build_workbook([_valid_row('1001')])})
                expect(page.locator('#fulfillmentAcceptance')).to_contain_text('Проверяем сохранение',timeout=10000)
                expect(page.locator('#fulfillmentUploadButton')).to_be_disabled()
                assert page.locator('#fulfillmentAcceptance .ff-operation-check').count()==0
                assert mutations==['upload']
                # Reload retains the same uncertain request and sends only GET.
                open_form(page,base)
                expect(page.locator('#fulfillmentAcceptance')).to_contain_text('Проверяем сохранение')
                assert mutations==['upload']
                foreign=False
                page.locator('#fulfillmentAcceptance').get_by_role('button',name='Проверить статус').click()
                expect(page.locator('#fulfillmentAcceptance .ff-operation-check')).to_be_visible(timeout=10000)
                expect(page.locator('#fulfillmentAcceptance')).to_contain_text('Документ сохранён.')
                expect(page.locator('#fulfillmentAcceptance')).to_contain_text('Ожидает обработки')
                assert page.locator('#fulfillmentAcceptance').get_by_text('Обработано',exact=True).count()==0
                expect(page.locator('#fulfillmentUploadButton')).to_be_enabled()
                page.locator('#fulfillmentAcceptance').get_by_role('button',name='Закрыть').focus()
                page.keyboard.press('Enter');expect(page.locator('#fulfillmentAcceptance')).to_be_hidden()
                with receipt.readonly(rt.db_path) as conn:
                    upload_id=conn.execute('SELECT upload_id FROM '+receipt.TABLE).fetchone()[0]
                def deletion(route):
                    mutations.append('delete');route.fetch();route.abort('failed')
                page.route('**/fulfillment-services/uploads/*?request_id=*',deletion)
                page.locator('[data-fulfillment-delete]').first.click()
                page.locator('[data-fulfillment-delete-confirm]').first.click()
                expect(page.locator('#fulfillmentAcceptance .ff-operation-check')).to_be_visible(timeout=10000)
                expect(page.locator('#fulfillmentAcceptance')).to_contain_text('Удаление документа услуг ФФ')
                assert mutations==['upload','delete']
                expect(page.locator('#fulfillmentUploadsBody')).to_contain_text('Загруженных документов пока нет')
                # Safe host journal link is an explicit browser navigation callback.
                link=page.locator('#fulfillmentAcceptance').get_by_role('link',name='Журнал операций')
                expect(link).to_be_visible();assert not link.get_attribute('href').startswith(('javascript:','//'))
                accepted_id=page.locator('#fulfillmentAcceptance [data-ff-operation-receipt]').get_attribute('data-ff-operation-receipt')
                link.click()
                expect(page.get_by_role('heading',name='Журнал операций',exact=True)).to_be_visible()
                expect(page.locator('#journal-detail [data-ff-operation-receipt]')).to_have_attribute('data-ff-operation-receipt',accepted_id)
                assert mutations==['upload','delete']
                open_form(page,base)
                with receipt.readonly(rt.db_path) as conn:
                    assert conn.execute('SELECT count(*) FROM '+receipt.TABLE).fetchone()[0]==2
                    assert conn.execute('SELECT count(*) FROM '+receipt.COMPLETIONS).fetchone()[0]==0
                old=_wb_supply_row('4001',accepted_quantity=10,quantity_added=10,cost_total=0)
                old.update(fact_date='2026-06-30',supply_date='2026-06-30')
                rt.save_wb_supply_rows(rows=[old],warehouses=[],synced_at='2026-07-10T10:00:00Z')
                choose_workbook(page, {'name':'outside-window.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':_build_workbook([_valid_row('4001')])})
                expect(page.locator('#fulfillmentAcceptance')).to_contain_text('Обработка не требуется',timeout=10000)
                assert page.locator('#fulfillmentAcceptance').get_by_text('Обработано',exact=True).count()==0
                assert mutations==['upload','delete','upload']
                choose_workbook(page, {'name':'diagnostic.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':_build_workbook([_valid_row('missing')])})
                expect(page.locator('#fulfillmentServicesMessage')).to_contain_text('Документ не принят.',timeout=10000)
                expect(page.locator('#fulfillmentAcceptance')).to_be_hidden()
                assert page.locator('#fulfillmentAcceptance .ff-operation-check').count()==0
                expect(page.locator('#fulfillmentUploadButton')).to_be_enabled()
                assert mutations==['upload','delete','upload','upload']
                browser.close()
            print('Chromium native upload/delete lost response, foreign request no green, reload same GET, pending receipt, keyboard close/journal: OK')
        finally:stop(server,thread)
    test_chartsheet_diagnostic_retry()


if __name__=='__main__':main()
