"""Exact operator catalog receipts with real native consumers; temporary data only."""
from contextlib import contextmanager
from datetime import date,timedelta
from pathlib import Path
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from openpyxl import load_workbook
import json,sqlite3,sys,tempfile,unittest
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from apps.nomenclature_activation_intents_smoke import fixture,ordinary,Interrupted
from apps.ff_pool_dense_fbs_smoke import _sku,NOW
from apps.nomenclature_barcode_smoke import FakeBarcodeSource
from apps.business_data_heavy_producers_smoke import busy
from apps.wb_finance_weekly_smoke import _seed_canonical_cost
from apps.wb_finance_daily_smoke import _rows,NOW as FINANCE_NOW
from packages.application import operator_nomenclature as op
from packages.application.business_data_heavy_admission import heavy_admitted
from packages.application.supplier_shipments import SupplierShipmentsBlock
from packages.application.registry_upload_http_entrypoint import RegistryUploadHttpEntrypoint as Entry
from packages.application.wb_finance_weekly import WbFinanceWeeklyBlock

class Tests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.runtime=fixture(Path(self.temp.name),facilities=1)
        self.provider=FakeBarcodeSource()
        self.supplier=SupplierShipmentsBlock(runtime=self.runtime,barcode_source=self.provider,timestamp_factory=lambda:NOW)
        self.supplier.list_sku_groups()
        self.supplier.create_sku_group({"group_key":"clear","label":"Clean"})
        self.entry=Entry.__new__(Entry);self.entry.runtime=self.runtime;self.entry.supplier_shipments_block=self.supplier
        self.counter=0
    def identity(self):
        self.counter+=1;return 'opsku_'+format(self.counter,'032x')
    def payload(self,**extra):
        return {'is_active':False,'nm_id':101,'nomenclature_name':'Synthetic iPhone 16',
            'product_type':'clear','match_key':'clear|test','barcode':'4600000000101',
            'purchase_price_yuan':1,'_operator_request_id':self.identity(),**extra}
    def read(self,result):
        return op.read(self.runtime.db_path,result['acceptance']['operation_id'],actor='tester')
    def create(self,**extra):
        return self.entry.handle_nomenclature_create_request(self.payload(**extra),actor='tester')
    def finance(self):
        block=WbFinanceWeeklyBlock(self.runtime.runtime_dir,seller_id='seller-1',now_factory=lambda:FINANCE_NOW)
        block.ensure_schema();return block
    def test_busy_source_exact_id_actor_and_unknown_do_not_resubmit(self):
        payload=self.payload()
        with busy(self.runtime.runtime_dir),patch.object(WbFinanceWeeklyBlock,'recalculate_stale_cost_weeks',side_effect=AssertionError('HTTP heavy')):
            first=self.entry.handle_nomenclature_create_request(payload,actor='tester')
            second=self.entry.handle_nomenclature_create_request(payload,actor='tester')
        self.assertEqual(first['acceptance'],second['acceptance'])
        self.assertEqual(len(self.runtime.list_nomenclature_items()),1)
        self.assertEqual(first['acceptance']['state'],'processing')
        self.assertIsNone(op.read(self.runtime.db_path,first['acceptance']['operation_id'],actor='foreign'))
        with self.assertRaisesRegex(ValueError,'identity_conflict'):
            self.entry.handle_nomenclature_create_request({**payload,'purchase_price_yuan':2},actor='tester')
        with self.assertRaisesRegex(ValueError,'identity_conflict'):
            self.entry.handle_nomenclature_create_request(payload,actor='foreign')
    def test_source_receipt_immutable_and_append_only(self):
        saved=self.create();identity=saved['acceptance']['operation_id']
        with sqlite3.connect(self.runtime.db_path) as conn:
            for column,value in [('actor','foreign'),('source_digest','sha256:foreign'),('result_json','{}')]:
                with self.assertRaisesRegex(sqlite3.IntegrityError,'immutable'):
                    conn.execute(f'UPDATE {op.TABLE} SET {column}=? WHERE operation_id=?',(value,identity))
            with self.assertRaisesRegex(sqlite3.IntegrityError,'append-only'):
                conn.execute(f'DELETE FROM {op.TABLE} WHERE operation_id=?',(identity,))
        self.assertEqual(self.read(saved)['operation_id'],identity)
    def test_source_receipt_atomic_and_lost_native_response_reconciled(self):
        original=op.record_saved
        with patch.object(op,'record_saved',side_effect=lambda *a,**kw:(_ for _ in ()).throw(RuntimeError('before_commit'))):
            with self.assertRaisesRegex(RuntimeError,'before_commit'):self.create()
        self.assertEqual(self.runtime.list_nomenclature_items(),[])
        native=self.runtime.save_nomenclature_item
        def lose(*args,**kwargs):
            native(*args,**kwargs);raise RuntimeError('after_source_commit')
        with patch.object(self.runtime,'save_nomenclature_item',side_effect=lose):saved=self.create()
        self.assertTrue(saved['acceptance']['durable_saved']);self.assertEqual(len(self.runtime.list_nomenclature_items()),1)
    def test_two_versions_cas_and_exact_supersession(self):
        first=self.create();item=first['item'];before=self.runtime.load_nomenclature_item(item['item_id'])
        second=self.entry.handle_nomenclature_patch_request(item['item_id'],{'comment':'new',
            '_operator_request_id':self.identity(),'_operator_expected_revision':before['operator_source_revision']},actor='tester')
        with self.assertRaisesRegex(ValueError,'version_changed'):
            self.entry.handle_nomenclature_patch_request(item['item_id'],{'comment':'stale','_operator_request_id':self.identity(),
                '_operator_expected_revision':before['operator_source_revision']},actor='tester')
        self.assertEqual(self.read(first)['reason_code'],'source_version_superseded')
        self.finance().recalculate_stale_cost_weeks()
        self.assertEqual(self.read(second)['state'],'completed')
        self.assertEqual(self.read(first)['state'],'needs_attention')
    def test_actual_finance_no_change_and_price_does_not_reprice_history(self):
        first=self.create();block=self.finance()
        result=block.recalculate_stale_cost_weeks()
        receipt=self.read(first)
        self.assertEqual(receipt['state'],'completed')
        proof=receipt['processing_receipt']['cost']
        self.assertEqual(proof['outcome'],'derived_no_change');self.assertTrue(proof['exact_revision_acked'])
        self.assertEqual(proof['sources'][0]['revision'],self.runtime.load_nomenclature_item(first['item']['item_id'])['operator_source_revision'])
        self.assertEqual(proof['catalog_price_policy'],'supplier_reference_not_historical_cost')
        source=self.runtime.load_nomenclature_item(first['item']['item_id'])
        changed=self.entry.handle_nomenclature_patch_request(source['item_id'],{'purchase_price_yuan':900,'_operator_request_id':self.identity()},actor='tester')
        self.assertEqual(self.read(changed)['state'],'processing')
        block.recalculate_stale_cost_weeks();self.assertEqual(self.read(changed)['state'],'completed')
    def test_actual_finance_applied_mapping_and_crash_before_ack_recovers(self):
        first=self.create(vendor_code='VC101',is_active=False)
        seed=self.runtime.runtime_dir/'seed.sqlite3';_seed_canonical_cost(seed)
        with sqlite3.connect(self.runtime.db_path) as conn,sqlite3.connect(seed) as source:
            for table in ('sheet_vitrina_v1_warehouse_functional_cutovers','sheet_vitrina_v1_warehouse_wb_daily_cost'):
                sql=source.execute('SELECT sql FROM sqlite_master WHERE name=?',(table,)).fetchone()[0]
                conn.execute(sql.replace('CREATE TABLE ','CREATE TABLE IF NOT EXISTS ',1))
                rows=source.execute('SELECT * FROM '+table).fetchall()
                conn.executemany('INSERT INTO '+table+' VALUES ('+','.join('?' for _ in rows[0])+')',rows)
        block=self.finance();day=date(2026,9,28)
        row=dict(_rows(day)[0],nmId=0,vendorCode='VC101',sku='',saleDt='2026-07-01')
        block.ingest_week(day,day+timedelta(days=6),[row]);block.recalculate_stale_cost_weeks()
        changed=self.entry.handle_nomenclature_patch_request(first['item']['item_id'],{'vendor_code':'VC999','_operator_request_id':self.identity()},actor='tester')
        plan=block.plan_stale_cost_weeks();self.assertEqual(plan['stale_week_count'],1)
        with patch.object(op,'acknowledge_finance',side_effect=Interrupted):
            with self.assertRaises(Interrupted):block.apply_stale_cost_weeks(expected_fingerprint=plan['fingerprint'])
        self.assertEqual(self.read(changed)['state'],'processing')
        recovered=block.recalculate_stale_cost_weeks();self.assertEqual(recovered['status'],'already_current')
        self.assertEqual(self.read(changed)['state'],'completed')
        restored=self.entry.handle_nomenclature_patch_request(first['item']['item_id'],{'vendor_code':'VC101','_operator_request_id':self.identity()},actor='tester')
        result=block.recalculate_stale_cost_weeks();self.assertEqual(result['status'],'applied')
        self.assertEqual(self.read(restored)['processing_receipt']['cost']['outcome'],'native_cost_evaluated')
        with sqlite3.connect(self.runtime.db_path) as conn:
            before=json.loads(conn.execute('SELECT metrics_json FROM wb_finance_weekly_aggregates').fetchone()[0])['cogs']
        price=self.entry.handle_nomenclature_patch_request(first['item']['item_id'],{'purchase_price_yuan':900,'_operator_request_id':self.identity()},actor='tester')
        block.recalculate_stale_cost_weeks()
        with sqlite3.connect(self.runtime.db_path) as conn:
            after=json.loads(conn.execute('SELECT metrics_json FROM wb_finance_weekly_aggregates').fetchone()[0])['cogs']
        self.assertEqual(before,after);self.assertEqual(self.read(price)['state'],'completed')
    def test_empty_barcode_before_barcode_only_finance_mapping_has_no_deadlock(self):
        seed=self.runtime.runtime_dir/'seed.sqlite3';_seed_canonical_cost(seed)
        with sqlite3.connect(self.runtime.db_path) as conn,sqlite3.connect(seed) as source:
            for table in ('sheet_vitrina_v1_warehouse_functional_cutovers','sheet_vitrina_v1_warehouse_wb_daily_cost'):
                sql=source.execute('SELECT sql FROM sqlite_master WHERE name=?',(table,)).fetchone()[0]
                conn.execute(sql.replace('CREATE TABLE ','CREATE TABLE IF NOT EXISTS ',1))
                rows=source.execute('SELECT * FROM '+table).fetchall()
                conn.executemany('INSERT INTO '+table+' VALUES ('+','.join('?' for _ in rows[0])+')',rows)
        block=self.finance();day=date(2026,9,28)
        row=dict(_rows(day)[0],nmId=0,vendorCode='',sku='4600000000101',saleDt='2026-07-01')
        block.ingest_week(day,day+timedelta(days=6),[row])
        self.provider.mapping={101:['4600000000101']}
        first=self.create(is_active=True,barcode='')
        ordinary(self.runtime)
        initial=block.recalculate_stale_cost_weeks()
        self.assertTrue(self.read(first)['processing_receipt']['cost']['exact_revision_acked'])
        self.assertEqual(self.read(first)['state'],'processing')
        with heavy_admitted(self.runtime.runtime_dir,operation='fixture'):op.drain_external(self.runtime,block=self.supplier)
        result=block.recalculate_stale_cost_weeks()
        self.assertEqual(self.read(first)['state'],'completed')
        self.assertEqual(result['post_verify_stale_week_count'],0)
        self.assertEqual(self.provider.calls,[[101]])
    def test_native_dense_publication_and_finance_both_required(self):
        first=self.create(is_active=True)
        item=self.runtime.load_nomenclature_item(first['item']['item_id'])
        self.assertFalse(item['is_active']);self.assertEqual(item['activation_status'],'pending')
        block=self.finance();block.recalculate_stale_cost_weeks()
        self.assertEqual(self.read(first)['state'],'processing')
        ordinary(self.runtime)
        self.assertTrue(self.runtime.load_nomenclature_item(item['item_id'])['is_active'])
        receipt=self.read(first);self.assertEqual(receipt['state'],'processing')
        self.assertTrue(receipt['processing_receipt']['activation'][item['item_id']]['dense_intent_id'])
        block.recalculate_stale_cost_weeks();self.assertEqual(self.read(first)['state'],'completed')
        ordinary(self.runtime);self.assertEqual(self.read(first)['state'],'completed')
    def test_auto_empty_barcode_native_chain_and_race_before_child(self):
        self.provider.mapping={101:['4600000000101']}
        with busy(self.runtime.runtime_dir):first=self.create(is_active=True,barcode='')
        self.assertEqual(self.provider.calls,[])
        self.assertEqual(first['barcode_sync']['status'],'pending')
        ordinary(self.runtime)
        block=self.finance();block.recalculate_stale_cost_weeks()
        self.assertEqual(self.read(first)['reason_code'],'native_barcode_pending')
        with heavy_admitted(self.runtime.runtime_dir,operation='fixture'):op.drain_external(self.runtime,block=self.supplier)
        self.assertEqual(self.provider.calls,[[101]])
        self.assertEqual(self.runtime.load_nomenclature_item(first['item']['item_id'])['barcode'],'4600000000101')
        self.assertEqual(self.read(first)['state'],'processing')
        block.recalculate_stale_cost_weeks();self.assertEqual(self.read(first)['state'],'completed')
        second=self.create(nm_id=102,match_key='clear|second',barcode='')
        block.recalculate_stale_cost_weeks()
        native=self.runtime.save_nomenclature_item
        def edit_before_child(item,**kw):
            original=self.runtime.load_nomenclature_item(item['item_id'])
            self.entry.handle_nomenclature_patch_request(item['item_id'],{'barcode':'4600000009999',
                'nm_id':103,'_operator_request_id':self.identity()},actor='tester')
            return native(item,**kw)
        # Apply explicit edit after WB snapshot but before child's writer CAS.
        with heavy_admitted(self.runtime.runtime_dir,operation='fixture'):
            with patch.object(self.runtime,'save_nomenclature_item',side_effect=edit_before_child):
                # Avoid recursive edit hook for the explicit operator write.
                def one_shot(item,**kw):
                    with patch.object(self.runtime,'save_nomenclature_item',side_effect=native):
                        return edit_before_child(item,**kw)
                with patch.object(self.runtime,'save_nomenclature_item',side_effect=one_shot):op.drain_external(self.runtime,block=self.supplier)
        current=self.runtime.load_nomenclature_item(second['item']['item_id'])
        self.assertEqual(current['barcode'],'4600000009999',self.read(second));self.assertEqual(current['nm_id'],103)
        self.assertEqual(self.read(second)['state'],'needs_attention')
    def test_auto_crash_after_child_recovers_without_repeating_provider_or_source(self):
        self.provider.mapping={101:['4600000000101']}
        first=self.create(barcode='');block=self.finance();block.recalculate_stale_cost_weeks()
        native=self.runtime.save_nomenclature_item
        def lose(*a,**kw):native(*a,**kw);raise Interrupted()
        with heavy_admitted(self.runtime.runtime_dir,operation='fixture'):
            with patch.object(self.runtime,'save_nomenclature_item',side_effect=lose):
                with self.assertRaises(Interrupted):op.drain_external(self.runtime,block=self.supplier)
            op.drain_external(self.runtime,block=self.supplier)
        self.assertEqual(self.provider.calls,[[101]])
        self.assertEqual(self.read(first)['state'],'processing')
        block.recalculate_stale_cost_weeks();self.assertEqual(self.read(first)['state'],'completed')
    def test_finance_no_change_handoff_race_never_acks_foreign_current_version(self):
        first=self.create();block=self.finance();plan=block.plan_stale_cost_weeks()
        native=block._connect
        def raced():
            with sqlite3.connect(self.runtime.db_path) as conn:
                conn.execute("UPDATE sheet_vitrina_v1_nomenclature_items SET comment='foreign' WHERE item_id=?",(first['item']['item_id'],))
                conn.commit()
            return native()
        with patch.object(block,'_connect',side_effect=raced):
            with self.assertRaisesRegex(ValueError,'changed'):
                block.apply_stale_cost_weeks(expected_fingerprint=plan['fingerprint'])
        self.assertEqual(self.read(first)['processing_receipt']['cost'],{})
    def test_invalid_negative_and_pending_group_retirement_fail_before_receipt(self):
        with self.assertRaises(ValueError):self.create(purchase_price_yuan=-1)
        self.assertEqual(self.runtime.list_nomenclature_items(),[])
        first=self.create(is_active=True)
        with self.assertRaisesRegex(ValueError,'used'):
            self.entry.handle_sku_groups_delete_request('clear',actor='tester',request_id=self.identity())
        self.assertTrue(self.runtime.load_sku_group('clear')['is_active'])
    def test_import_exact_file_atomic_invalid_and_same_identity(self):
        first=self.create();self.finance().recalculate_stale_cost_weeks()
        body,filename,_type=self.supplier.export_nomenclature_xlsx();book=load_workbook(BytesIO(body))
        column=next(cell.column for cell in book.active[1] if cell.value=='Цена закупки, ¥')
        book.active.cell(2,column).value=2;output=BytesIO();book.save(output)
        identity=self.identity()
        source=self.runtime.load_nomenclature_item(first['item']['item_id'])
        accepted=self.entry.handle_nomenclature_import_request(output.getvalue(),uploaded_filename=filename,
            actor='tester',request_id=identity,expected_revision={source['item_id']:source['operator_source_revision']})
        again=self.entry.handle_nomenclature_import_request(output.getvalue(),uploaded_filename=filename,
            actor='tester',request_id=identity,expected_revision={source['item_id']:source['operator_source_revision']})
        self.assertEqual(accepted['acceptance'],again['acceptance']);self.assertEqual(len(self.runtime.list_nomenclature_items()),1)
        book.active.cell(2,column).value=-1;bad=BytesIO();book.save(bad)
        rejected=self.entry.handle_nomenclature_import_request(bad.getvalue(),uploaded_filename=filename,
            actor='tester',request_id=self.identity())
        self.assertEqual(rejected['status'],'error');self.assertNotIn('acceptance',rejected)
        self.assertEqual(self.runtime.load_nomenclature_item(source['item_id'])['purchase_price_yuan'],2)
    def test_concurrent_same_before_version_only_one_source_is_saved(self):
        first=self.create();source=self.runtime.load_nomenclature_item(first['item']['item_id'])
        barrier=Barrier(2);native=self.runtime.save_nomenclature_item
        def blocked(*a,**kw):barrier.wait(timeout=10);return native(*a,**kw)
        def write(identity,comment):
            try:return self.entry.handle_nomenclature_patch_request(source['item_id'],{'comment':comment,
                '_operator_request_id':identity,'_operator_expected_revision':source['operator_source_revision']},actor='tester')
            except ValueError as error:return str(error)
        with patch.object(self.runtime,'save_nomenclature_item',side_effect=blocked),ThreadPoolExecutor(max_workers=2) as pool:
            one=pool.submit(write,self.identity(),'first');two=pool.submit(write,self.identity(),'second')
            results=[one.result(),two.result()]
        self.assertEqual(sum(isinstance(result,dict) for result in results),1)
        self.assertTrue(any('version_changed' in result for result in results if isinstance(result,str)))
    def test_missing_completion_proof_does_not_mean_zero(self):
        first=self.create();self.finance().recalculate_stale_cost_weeks()
        identity=first['acceptance']['operation_id']
        with sqlite3.connect(self.runtime.db_path) as conn:
            proof=json.loads(conn.execute(f'SELECT cost_receipt_json FROM {op.TABLE} WHERE operation_id=?',(identity,)).fetchone()[0])
            proof.pop('post_verify_stale_week_count')
            conn.execute(f'UPDATE {op.TABLE} SET cost_receipt_json=? WHERE operation_id=?',(json.dumps(proof),identity));conn.commit()
        self.assertEqual(self.read(first)['state'],'needs_attention')
        self.assertEqual(self.read(first)['reason_code'],'native_finance_proof_missing')
    def test_read_only_unknown_does_not_bootstrap(self):
        empty=self.runtime.runtime_dir/'unknown.sqlite3'
        self.assertIsNone(op.read(empty,self.identity(),actor='tester'));self.assertFalse(empty.exists())
    def test_pending_export_preserves_enabled_and_exact_import_off_cancels(self):
        first=self.create(is_active=True)
        body,filename,_type=self.supplier.export_nomenclature_xlsx();book=load_workbook(BytesIO(body))
        column=next(cell.column for cell in book.active[1] if cell.value=='Включено')
        self.assertEqual(book.active.cell(2,column).value,'да')
        self.assertFalse(self.runtime.load_nomenclature_item(first['item']['item_id'])['is_active'])
        book.active.cell(2,column).value='нет';output=BytesIO();book.save(output)
        result=self.entry.handle_nomenclature_import_request(output.getvalue(),uploaded_filename=filename,
            actor='tester',request_id=self.identity())
        self.assertEqual(result['deactivated_count'],1)
        self.assertNotEqual(self.runtime.load_nomenclature_item(first['item']['item_id'])['activation_status'],'pending')
        ordinary(self.runtime)
        self.assertFalse(self.runtime.load_nomenclature_item(first['item']['item_id'])['is_active'])
    def test_groups_versioned_and_no_fake_warehouse_job(self):
        first=self.entry.handle_sku_groups_create_request({'group_key':'synthetic','label':'Synthetic',
            '_operator_request_id':self.identity()},actor='tester')
        self.assertEqual(self.read(first)['state'],'completed');self.assertEqual(self.read(first)['processing_receipt']['cost'],{})
        old=self.runtime.load_sku_group('synthetic')
        second=self.entry.handle_sku_groups_patch_request('synthetic',{'label':'Renamed','_operator_request_id':self.identity(),
            '_operator_expected_revision':old['operator_source_revision']},actor='tester')
        with self.assertRaisesRegex(ValueError,'version_changed'):
            self.entry.handle_sku_groups_patch_request('synthetic',{'label':'Stale','_operator_request_id':self.identity(),
            '_operator_expected_revision':old['operator_source_revision']},actor='tester')
        self.assertEqual(self.read(second)['state'],'completed')
    def test_external_read_pending_retry_snapshot_and_child_exact_completion(self):
        self.provider.cards=[{'nm_id':101,'vendor_code':'VC101','title':'iPhone 16','barcodes':['4600000000101']},
                             {'nm_id':102,'vendor_code':'VC102','title':'iPhone 17','barcodes':['4600000000102']}]
        payload={'limit':10,'max_pages':1,'_operator_request_id':self.identity()}
        with busy(self.runtime.runtime_dir):accepted=self.entry.handle_nomenclature_barcode_sync_request(payload,actor='tester')
        self.assertEqual(self.provider.card_calls,[]);self.assertEqual(self.runtime.list_nomenclature_items(),[])
        native=self.runtime.save_nomenclature_item;count=0
        def crash(*args,**kwargs):
            nonlocal count
            result=native(*args,**kwargs);count+=1
            if count==1:raise Interrupted()
            return result
        with heavy_admitted(self.runtime.runtime_dir,operation='fixture'):
            with patch.object(self.runtime,'save_nomenclature_item',side_effect=crash):
                with self.assertRaises(Interrupted):op.drain_external(self.runtime,block=self.supplier)
            self.assertEqual(len(self.runtime.list_nomenclature_items()),1)
            op.drain_external(self.runtime,block=self.supplier)
        self.assertEqual(len(self.provider.card_calls),1);self.assertEqual(len(self.runtime.list_nomenclature_items()),2)
        self.assertEqual(self.read(accepted)['state'],'processing')
        ordinary(self.runtime);self.finance().recalculate_stale_cost_weeks()
        self.assertEqual(self.read(accepted)['state'],'completed')
        with heavy_admitted(self.runtime.runtime_dir,operation='fixture'):self.assertEqual(op.drain_external(self.runtime,block=self.supplier),[])
    def test_barcode_explicit_task_stale_and_manual_no_change(self):
        first=self.create();identity=first['item']['item_id'];self.finance().recalculate_stale_cost_weeks()
        accepted=self.entry.handle_nomenclature_item_barcode_sync_request(identity,actor='tester',request_id=self.identity())
        with heavy_admitted(self.runtime.runtime_dir,operation='fixture'):op.drain_external(self.runtime,block=self.supplier)
        self.assertEqual(self.read(accepted)['state'],'completed');self.assertEqual(self.provider.calls,[])
        stale=self.entry.handle_nomenclature_item_barcode_sync_request(identity,actor='tester',request_id=self.identity())
        self.entry.handle_nomenclature_patch_request(identity,{'comment':'changed','_operator_request_id':self.identity()},actor='tester')
        with heavy_admitted(self.runtime.runtime_dir,operation='fixture'):op.drain_external(self.runtime,block=self.supplier)
        self.assertEqual(self.read(stale)['state'],'needs_attention');self.assertEqual(self.provider.calls,[])

if __name__=='__main__':unittest.main()
