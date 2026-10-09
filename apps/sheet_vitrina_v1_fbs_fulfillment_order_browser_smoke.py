#!/usr/bin/env python3
"""Browser contract smoke for the independent FBS fulfillment-order UI."""

from __future__ import annotations

from datetime import datetime, timezone
import json
from pathlib import Path
import socket
import sqlite3
import sys
from tempfile import TemporaryDirectory
import threading
from urllib.parse import parse_qs, urlparse
from unittest.mock import patch

from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from apps.fbs_fulfillment_order_supply_smoke import (  # noqa: E402
    INPUT_BUNDLE_FIXTURE,
    MOSCOW_ID,
    ORENBURG_ID,
    _seed_facilities,
    _seed_sales_history,
    _seed_fbs_demand,
    _seed_shipments,
)
from packages.adapters.registry_upload_http_entrypoint import (  # noqa: E402
    DEFAULT_SHEET_OPERATOR_UI_PATH,
    DEFAULT_SHEET_PLAN_PATH,
    DEFAULT_SHEET_REFRESH_PATH,
    DEFAULT_SHEET_STATUS_PATH,
    DEFAULT_UPLOAD_PATH,
    build_registry_upload_http_server,
)
from packages.application.registry_upload_db_backed_runtime import (  # noqa: E402
    RegistryUploadDbBackedRuntime,
)
from packages.application.registry_upload_http_entrypoint import (  # noqa: E402
    RegistryUploadHttpEntrypoint,
)
from packages.contracts.registry_upload_http_entrypoint import (  # noqa: E402
    RegistryUploadHttpEntrypointConfig,
)
from packages.application.fbs_fulfillment_order import FbsFulfillmentOrderBlock
from packages.application.wb_fbs_orders import OBSERVATIONS_TABLE  # noqa: E402


NOW = datetime(2026, 4, 18, 9, 0, tzinfo=timezone.utc)
NOW_TEXT = "2026-04-18T09:00:00Z"


def main() -> int:
    bundle = json.loads(INPUT_BUNDLE_FIXTURE.read_text(encoding="utf-8"))
    with TemporaryDirectory(prefix="fbs-fulfillment-browser-") as raw:
        runtime_dir = Path(raw) / "runtime"
        runtime = RegistryUploadDbBackedRuntime(runtime_dir=runtime_dir)
        runtime.ingest_bundle(bundle, activated_at=NOW_TEXT)
        active_nm_ids = [
            int(item.nm_id)
            for item in runtime.load_current_state().config_v2
            if item.enabled
        ]
        _seed_facilities(runtime, active_nm_ids)
        _seed_sales_history(runtime, active_nm_ids)
        _seed_fbs_demand(runtime, active_nm_ids)
        # The full-volume fixture must have the collector's order index. Without
        # it, the reader's per-order latest-revision lookup becomes quadratic;
        # a Linux CI runner spends >15 seconds in each readiness GET.
        with sqlite3.connect(runtime_dir / "fbs_observer" / "observations.sqlite3") as observer:
            plan = observer.execute(f"EXPLAIN QUERY PLAN SELECT MAX(observation_sequence) FROM {OBSERVATIONS_TABLE} WHERE order_id=?", (1,)).fetchall()
            assert any("wb_fbs_observations_by_order" in row[3] for row in plan), plan
        _seed_shipments(runtime, active_nm_ids)

        port = _reserve_free_port()
        entrypoint = RegistryUploadHttpEntrypoint(
            runtime_dir=runtime_dir,
            runtime=runtime,
            activated_at_factory=lambda: NOW_TEXT,
            now_factory=lambda: NOW,
        )
        config = RegistryUploadHttpEntrypointConfig(
            host="127.0.0.1",
            port=port,
            upload_path=DEFAULT_UPLOAD_PATH,
            sheet_plan_path=DEFAULT_SHEET_PLAN_PATH,
            sheet_refresh_path=DEFAULT_SHEET_REFRESH_PATH,
            sheet_status_path=DEFAULT_SHEET_STATUS_PATH,
            sheet_operator_ui_path=DEFAULT_SHEET_OPERATOR_UI_PATH,
            runtime_dir=runtime_dir,
        )
        server = build_registry_upload_http_server(config, entrypoint=entrypoint)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        page_errors: list[str] = []
        console_errors: list[str] = []
        fbs_http_errors: list[str] = []
        try:
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch()
                context = browser.new_context(viewport={"width": 1280, "height": 900})
                page = context.new_page()
                # Mark a status response only after its text has been consumed
                # and the application's promise continuation has run. This lets
                # us release a held old response and assert the resulting UI,
                # without a timing sleep that could miss the overwrite.
                page.add_init_script("""(() => {
                  const originalFetch = window.fetch;
                  window.__fbsStatusSettled = 0;
                  window.fetch = (...args) => originalFetch(...args).then(response => {
                    if (response.url.includes('fbs-fulfillment-order/status')) {
                      const originalText = response.text.bind(response);
                      response.text = async () => {
                        const text = await originalText();
                        setTimeout(() => { window.__fbsStatusSettled += 1; }, 0);
                        return text;
                      };
                    }
                    return response;
                  });
                })();""")
                held_status = []
                hold_next_status = {"enabled": False}
                forced_status_blocker = {"enabled": False}
                def intercept_status(route):
                    if not hold_next_status["enabled"] and not forced_status_blocker["enabled"]:
                        route.continue_()
                        return
                    hold = hold_next_status["enabled"]
                    hold_next_status["enabled"] = False
                    response = route.fetch()
                    try:
                        payload = response.json()
                    finally:
                        response.dispose()
                    if forced_status_blocker["enabled"]:
                        for item in payload["facilities"]:
                            item.update(source_blocker="Нет полного официального снимка", calculation_enabled=False, blockers=["Нет полного официального снимка"])
                    if not hold:
                        route.fulfill(json=payload)
                        return
                    held_status.append((route, payload))
                    page.evaluate("document.documentElement.dataset.fbsHeldStatus = String(Number(document.documentElement.dataset.fbsHeldStatus || 0) + 1)")
                page.route("**/fbs-fulfillment-order/status*", intercept_status)
                def hold_status_request():
                    old_count = page.evaluate("Number(document.documentElement.dataset.fbsHeldStatus || 0)")
                    hold_next_status["enabled"] = True
                    page.locator('[data-supply-section-button="fbs-fulfillment"]').click()
                    page.wait_for_function("count => Number(document.documentElement.dataset.fbsHeldStatus || 0) > count", arg=old_count)
                def release_status_request():
                    previous = page.evaluate("window.__fbsStatusSettled")
                    route, payload = held_status.pop(0)
                    cancelled = route.request.failure is not None
                    route.fulfill(json=payload)
                    if not cancelled:
                        page.wait_for_function("count => window.__fbsStatusSettled > count", arg=previous)
                held_calculations = []
                calculation_payloads = []
                fbs_requests = []
                page.on("request", lambda request: fbs_requests.append(request.url) if "/fbs-fulfillment-order/" in request.url else None)
                hold_next_calculation = {"enabled": False}
                def intercept_calculation(route):
                    calculation_payloads.append(route.request.post_data_json)
                    if not hold_next_calculation["enabled"]:
                        route.continue_()
                        return
                    hold_next_calculation["enabled"] = False
                    held_calculations.append(route)
                    page.evaluate("document.documentElement.dataset.fbsHeldCalculation = String(Number(document.documentElement.dataset.fbsHeldCalculation || 0) + 1)")
                page.route("**/fbs-fulfillment-order/calculate", intercept_calculation)
                def hold_calculation_request():
                    old_count = page.evaluate("Number(document.documentElement.dataset.fbsHeldCalculation || 0)")
                    hold_next_calculation["enabled"] = True
                    page.locator("#fbsFulfillmentCalculateButton").click()
                    page.wait_for_function("count => Number(document.documentElement.dataset.fbsHeldCalculation || 0) > count", arg=old_count)
                page.on("pageerror", lambda error: page_errors.append(str(error)))
                page.on(
                    "console",
                    lambda message: console_errors.append(message.text)
                    if message.type == "error"
                    and not message.text.startswith("Failed to load resource:")
                    else None,
                )
                page.on(
                    "response",
                    lambda response: fbs_http_errors.append(
                        f"{response.status} {response.url}"
                    )
                    if "fbs-fulfillment-order" in response.url
                    and response.status >= 400
                    and response.headers.get("x-wbc-fixture-expected-failure") != "1"
                    else None,
                )
                page.goto(
                    f"http://127.0.0.1:{port}{DEFAULT_SHEET_OPERATOR_UI_PATH}"
                    "?embedded_tab=factory-order",
                    wait_until="domcontentloaded",
                )

                fbs_panel = page.locator(
                    '[data-supply-section-panel="fbs-fulfillment"]'
                )
                legacy_panel = page.locator('[data-supply-section-panel="factory"]')
                expect(fbs_panel).to_be_visible()
                expect(legacy_panel).to_be_hidden()
                expect(page.locator("#legacyFactoryOrderDetails")).not_to_have_attribute(
                    "open", ""
                )
                expect(page.locator("#fbsHistoryModeLastN")).to_be_checked()
                expect(page.locator("#fbsInboundScope")).to_have_value(
                    "selected_facility"
                )
                expect(page.locator("#fbsInboundScopeHelp")).to_contain_text(
                    "нельзя одновременно считать распределённым"
                )
                expect(page.locator("#fbsReadinessDetails")).not_to_have_attribute(
                    "open", ""
                )
                assert page.locator("#fbsFulfillmentResultBody").count() == 0
                assert fbs_panel.evaluate(
                    "element => element.scrollWidth <= element.clientWidth + 1"
                )
                expect(page.locator("#fbsSalesAvgPeriodDays")).to_have_value("14")
                expect(page.locator("#fbsSalesDateFrom")).to_be_disabled()
                expect(page.locator("#fbsSalesDateTo")).to_be_disabled()

                facility = page.locator("#fbsTargetFacility")
                expect(facility.locator("option")).to_have_count(3, timeout=15000)
                expect(facility).to_have_value(MOSCOW_ID)
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_enabled()
                expect(page.locator("#fbsReadinessPhysical")).not_to_have_text("-")
                expect(page.locator("#fbsReadinessReserved")).not_to_have_text("-")
                expect(page.locator("#fbsReadinessAvailable")).not_to_have_text("-")
                expect(page.locator("#fbsHistoryCoverage")).to_contain_text(
                    "2026-04-02 — 2026-04-17"
                )

                # The full eligible catalog starts selected. Category controls
                # affect the calculation scope; the server certifies subset readiness.
                expect(page.locator("#fbsSkuSelectionSummary")).to_have_text(f"SKU к заказу ({len(active_nm_ids)}/{len(active_nm_ids)})")
                assert page.locator(".fbs-settings-card").evaluate("node => getComputedStyle(node).display === 'grid' && getComputedStyle(node).gap === '12px'")
                page.locator("#fbsSkuSelectionSummary").click()
                expect(page.locator("[data-fbs-sku]")).to_have_count(len(active_nm_ids))
                labels = page.locator(".fbs-sku-category").all_text_contents()
                assert labels == sorted(labels, key=lambda label: label.casefold())
                category = page.locator("[data-fbs-sku-category]").first
                category_ids = category.evaluate("node => {const ids=[];for(let label=node.parentElement.nextElementSibling;label&&!label.classList.contains('fbs-sku-category');label=label.nextElementSibling)ids.push(Number(label.querySelector('input').dataset.fbsSku));return ids}")
                assert len(category_ids) > 1
                requests_before_selection = len(fbs_requests)
                category.uncheck()
                assert page.locator("[data-fbs-sku]").evaluate_all("nodes => nodes.filter(n => !n.checked).map(n => Number(n.dataset.fbsSku))") == category_ids
                category.check()
                first_sku = page.locator("[data-fbs-sku]").first
                first_id = int(first_sku.get_attribute("data-fbs-sku"))
                first_sku.uncheck()
                assert category.evaluate("node => node.indeterminate")
                assert first_sku.evaluate("node => node === document.activeElement")
                page.locator("[data-fbs-sku-none]").click()
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_disabled()
                expect(page.locator("#fbsSkuSelectionNote")).to_contain_text("Не выбраны SKU")
                page.locator("[data-fbs-sku-all]").click()
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_enabled()
                assert not any("/calculate" in url for url in fbs_requests[requests_before_selection:])
                # Excluding a SKU must never bypass a global snapshot failure.
                forced_status_blocker["enabled"] = True
                hold_status_request()
                release_status_request()
                first_sku.uncheck()
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_disabled()
                expect(page.locator("#fbsReadinessBlockers")).to_contain_text("Нет полного официального снимка")
                first_sku.check()
                forced_status_blocker["enabled"] = False
                settled = page.evaluate("window.__fbsStatusSettled")
                page.locator('[data-supply-section-button="fbs-fulfillment"]').click()
                page.wait_for_function("count => window.__fbsStatusSettled > count", arg=settled)
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_enabled()
                page.locator("#fbsSkuSelectionSummary").click()

                facility.select_option(ORENBURG_ID)
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_enabled()
                expect(fbs_panel.locator(".source-pill").first).to_have_text("Спрос выбранного склада")
                facility.select_option(MOSCOW_ID)
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_enabled()

                page.locator("#fbsHistoryModeCustom").check()
                expect(page.locator("#fbsSalesAvgPeriodDays")).to_be_disabled()
                page.locator("#fbsSalesDateFrom").fill("2026-04-10")
                page.locator("#fbsSalesDateTo").fill("2026-04-12")
                hold_status_request()
                secondary_marker = page.evaluate("Number(document.documentElement.dataset.fbsHeldStatus || 0)")
                hold_next_status["enabled"] = True
                page.locator("#fbsFulfillmentCalculateButton").click()
                # Also hold the follow-up GET issued after the POST. A completed
                # recommendation must not wait for this unrelated readiness read.
                page.wait_for_function("count => Number(document.documentElement.dataset.fbsHeldStatus || 0) > count", arg=secondary_marker)
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text(
                    "Расчёт завершён", timeout=15000
                )
                expect(page.locator("#fbsDemandWindow")).to_contain_text(
                    "Произвольный период"
                )
                expect(page.locator("#fbsDemandWindow")).to_contain_text(
                    "2026-04-10 — 2026-04-12"
                )
                expect(page.locator("#fbsDemandDays")).to_contain_text("3 /")
                expect(page.locator("#fbsTotalQty")).not_to_have_text("-")
                expect(page.locator("#fbsHorizonDays")).to_have_text("89")
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_enabled()
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_enabled()
                expect(page.locator("#fbsPreviewNote")).to_contain_text("Рассчитано")
                expect(page.locator("#fbsResultInboundScope")).to_have_text(
                    "Только для выбранного ФФ"
                )
                expect(page.locator("#fbsResultInbound")).to_contain_text("35 шт.")
                result_qty = page.locator("#fbsTotalQty").inner_text()
                release_status_request()
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Расчёт завершён")
                release_status_request()
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Расчёт завершён")

                # A fresh readiness probe during a repeated POST includes the
                # previously saved result for the same form. It must not restore
                # that old preview or re-enable its Excel while the POST is held.
                hold_calculation_request()
                settled = page.evaluate("window.__fbsStatusSettled")
                page.locator('[data-supply-section-button="fbs-fulfillment"]').click()
                page.wait_for_function("count => window.__fbsStatusSettled > count", arg=settled)
                expect(page.locator("#fbsTotalQty")).to_have_text("—")
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_disabled()
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_disabled()
                held_calculations.pop(0).continue_()
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Расчёт завершён", timeout=15000)
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_enabled()
                expect(page.locator("#fbsTotalQty")).to_have_text(result_qty)
                expect(page.locator("#fbsDemandWindow")).to_contain_text("2026-04-10 — 2026-04-12")
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_enabled()
                settled = page.evaluate("window.__fbsStatusSettled")
                page.locator('[data-supply-section-button="fbs-fulfillment"]').click()
                page.wait_for_function("count => window.__fbsStatusSettled > count", arg=settled)
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Расчёт завершён")

                with page.expect_download(timeout=15000) as download_info:
                    page.locator("#fbsFulfillmentDownloadButton").click()
                download = download_info.value
                assert download.suggested_filename.endswith(".xlsx")
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text(
                    "FBS-рекомендация скачана"
                )

                hold_status_request()
                page.locator("#fbsInboundScope").select_option("all_active")
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text(
                    "Параметры изменены"
                )
                expect(page.locator("#fbsTotalQty")).to_have_text("—")
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_disabled()
                release_status_request()
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Параметры изменены")
                expect(page.locator("#fbsTotalQty")).to_have_text("—")
                page.locator("#fbsFulfillmentCalculateButton").click()
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text(
                    "Расчёт завершён", timeout=15000
                )
                expect(page.locator("#fbsResultInboundScope")).to_have_text(
                    "Все активные заказы фабрике"
                )
                expect(page.locator("#fbsResultInbound")).to_contain_text("535 шт.")

                # A failing older POST must not overwrite the newer form's
                # changed-parameters warning.
                hold_calculation_request()
                page.locator("#fbsOrderBatchQty").fill("251")
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Параметры изменены")
                held_calculations.pop(0).fulfill(status=500, headers={"X-WBC-Fixture-Expected-Failure": "1"}, json={"error": "delayed_fixture_failure"})
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_have_text("Рассчитать заказ")
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Параметры изменены")
                expect(page.locator("#fbsTotalQty")).to_have_text("—")
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_disabled()

                page.locator("#fbsHistoryModeLastN").check()
                expect(page.locator("#fbsSalesAvgPeriodDays")).to_be_enabled()
                expect(page.locator("#fbsSalesDateFrom")).to_be_disabled()
                expect(page.locator("#fbsSalesDateTo")).to_be_disabled()

                # Changing SKU scope invalidates the existing Excel and cannot
                # be undone by an older readiness response.
                page.locator("#fbsSkuSelectionSummary").click()
                hold_status_request()
                for stale_facility in held_status[0][1]["facilities"]:
                    stale_facility["available"] = 987654321
                first_sku.uncheck()
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_disabled()
                expect(page.locator("#fbsTotalQty")).to_have_text("—")
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_disabled()
                first_sku.evaluate("node => window.oldSkuCheckbox = node")
                first_sku.focus()
                release_status_request()
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_enabled()
                expect(page.locator("#fbsReadinessAvailable")).not_to_have_text("987654321")
                assert first_sku.evaluate("node => node === window.oldSkuCheckbox && node === document.activeElement")
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_disabled()
                # Removed catalog IDs may remain in local preferences, but must
                # not be sent to the strict current-catalog API.
                page.evaluate("key => {const ids=JSON.parse(localStorage.getItem(key));ids.push(999999999);localStorage.setItem(key,JSON.stringify(ids));localStorage.setItem('wbc.stock-monitor.hidden-skus.v1',JSON.stringify([123]));}", "wbc.fbs-fulfillment.excluded-skus.v1")
                page.reload(wait_until="domcontentloaded")
                expect(page.locator("#fbsSkuSelectionSummary")).to_have_text(f"SKU к заказу ({len(active_nm_ids)-1}/{len(active_nm_ids)})", timeout=15000)
                # Catalog rendering precedes the persisted selection's scoped
                # readiness GET. Finish that bootstrap before arming a held
                # metadata response, so it belongs to the explicit activation.
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_enabled()
                page.locator("#fbsSkuSelectionSummary").click()
                first_sku = page.locator(f'[data-fbs-sku="{first_id}"]')
                expect(first_sku).not_to_be_checked()
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_disabled()
                assert page.evaluate("localStorage.getItem('wbc.stock-monitor.hidden-skus.v1')") == "[123]"
                # A newly added eligible SKU is selected implicitly, and a real
                # catalog metadata change restores focus by stable identity.
                hold_status_request()
                route, incoming = held_status[0]
                incoming["sku_catalog"].append({"nm_id": 999999998, "name": "Новый SKU", "category": "No Frame Clean"})
                first_sku.focus()
                release_status_request()
                expect(page.locator('[data-fbs-sku="999999998"]')).to_be_checked()
                assert first_sku.evaluate("node => node === document.activeElement")
                settled = page.evaluate("window.__fbsStatusSettled")
                page.locator('[data-supply-section-button="fbs-fulfillment"]').click()
                page.wait_for_function("count => window.__fbsStatusSettled > count", arg=settled)
                expect(page.locator('[data-fbs-sku="999999998"]')).to_have_count(0)
                page.locator("#fbsFulfillmentCalculateButton").click()
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Расчёт завершён", timeout=15000)
                assert calculation_payloads[-1]["excluded_nm_ids"] == [first_id]
                saved = runtime.load_fbs_fulfillment_order_result_state()
                assert saved["settings"]["excluded_nm_ids"] == [first_id]
                assert first_id not in {row["nm_id"] for row in saved["rows"]}
                assert len(saved["rows"]) == len(active_nm_ids) - 1
                # Another client's calculation replaces the latest snapshot.
                # This tab must export the immutable calculation it displays.
                block = FbsFulfillmentOrderBlock(runtime=runtime, now_factory=lambda: NOW, timestamp_factory=lambda: NOW_TEXT)
                other = block.calculate({**saved["settings"], "excluded_nm_ids": []})
                assert other.calculation_id != saved["calculation_id"]
                assert first_id in {row.nm_id for row in other.rows}
                with page.expect_download(timeout=15000) as selected_download:
                    page.locator("#fbsFulfillmentDownloadButton").click()
                assert any("/recommendation.xlsx?" in url and parse_qs(urlparse(url).query).get("calculation_id") == [saved["calculation_id"]] for url in fbs_requests)
                selected_path = Path(raw) / "selected-recommendation.xlsx"
                selected_download.value.save_as(selected_path)
                from openpyxl import load_workbook
                workbook = load_workbook(selected_path, read_only=True, data_only=True)
                exported = {int(row[0]) for row in workbook.active.iter_rows(min_row=2, values_only=True) if str(row[0] or "").isdigit()}
                workbook.close()
                assert first_id not in exported
                assert exported == {row["nm_id"] for row in saved["rows"] if row["recommended_order_qty"] is not None}
                # A successful POST sent for an earlier SKU choice must never
                # restore its saved recommendation after another checkbox edit.
                hold_calculation_request()
                second_sku = page.locator("[data-fbs-sku]:checked").first
                second_id = int(second_sku.get_attribute("data-fbs-sku"))
                second_sku.uncheck()
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_disabled()
                held_calculations.pop(0).continue_()
                expect(page.locator("#fbsFulfillmentCalculateButton")).to_have_text("Рассчитать заказ")
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Параметры изменены")
                expect(page.locator("#fbsTotalQty")).to_have_text("—")
                settled = page.evaluate("window.__fbsStatusSettled")
                page.locator('[data-supply-section-button="fbs-fulfillment"]').click()
                page.wait_for_function("count => window.__fbsStatusSettled > count", arg=settled)
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_disabled()
                expect(page.locator("#fbsTotalQty")).to_have_text("—")
                page.locator("#fbsFulfillmentCalculateButton").click()
                expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Расчёт завершён", timeout=15000)
                assert calculation_payloads[-1]["excluded_nm_ids"] == sorted([first_id, second_id])
                assert {first_id, second_id}.isdisjoint(row["nm_id"] for row in runtime.load_fbs_fulfillment_order_result_state()["rows"])

                # A new eligible SKU lies outside the last complete official
                # generation. ALL is blocked; excluding it is certified by the
                # real scoped API while its full catalog stays available.
                eligible = block._load_active_skus()
                new_id = 999999999998
                with patch.object(FbsFulfillmentOrderBlock, "_load_active_skus", return_value=eligible + [(new_id, "Новый SKU")]):
                    page.locator("[data-fbs-sku-all]").click()
                    expect(page.locator(f'[data-fbs-sku="{new_id}"]')).to_be_checked()
                    expect(page.locator("#fbsReadinessBlockers")).to_contain_text("Нет полного официального снимка", timeout=15000)
                    expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_disabled()
                    requests_before_none = len(fbs_requests)
                    page.locator("[data-fbs-sku-none]").click()
                    # Wait past the debounce deadline to prove empty scope
                    # cannot request readiness or a calculation.
                    page.wait_for_timeout(450)
                    assert len(fbs_requests) == requests_before_none
                    page.locator("[data-fbs-sku-all]").click()
                    expect(page.locator("#fbsReadinessBlockers")).to_contain_text("Нет полного официального снимка", timeout=15000)
                    page.locator(f'[data-fbs-sku="{new_id}"]').uncheck()
                    expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_disabled()
                    expect(page.locator("#fbsFulfillmentCalculateButton")).to_be_enabled(timeout=15000)
                    expect(page.locator("[data-fbs-sku]")).to_have_count(len(active_nm_ids) + 1)
                    expect(page.locator("#fbsSkuSelectionSummary")).to_have_text(f"SKU к заказу ({len(active_nm_ids)}/{len(active_nm_ids)+1})")
                    assert any(json.loads(parse_qs(urlparse(url).query)["excluded_nm_ids"][0]) == [new_id] for url in fbs_requests if "/status?" in url)
                    page.locator("#fbsFulfillmentCalculateButton").click()
                    expect(page.locator("#fbsFulfillmentMessage")).to_contain_text("Расчёт завершён", timeout=15000)
                    assert calculation_payloads[-1]["excluded_nm_ids"] == [new_id]
                    assert {row["nm_id"] for row in runtime.load_fbs_fulfillment_order_result_state()["rows"]} == set(active_nm_ids)

                page.set_viewport_size({"width": 390, "height": 844})
                assert fbs_panel.evaluate(
                    "element => element.scrollWidth <= element.clientWidth + 1"
                )
                expect(page.locator("#fbsInboundScope")).to_be_visible()
                expect(page.locator("#fbsFulfillmentDownloadButton")).to_be_visible()

                context.close()
                browser.close()
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)

        if page_errors:
            raise AssertionError(f"browser page errors: {page_errors}")
        if console_errors:
            raise AssertionError(f"browser console errors: {console_errors}")
        if fbs_http_errors:
            raise AssertionError(f"browser FBS HTTP errors: {fbs_http_errors}")
    print("sheet_vitrina_v1_fbs_fulfillment_order_browser_smoke: ok")
    return 0


def _reserve_free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


if __name__ == "__main__":
    raise SystemExit(main())
